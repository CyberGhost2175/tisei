#!/usr/bin/env python3
"""Собирает красивые .docx из markdown-файлов ТЗ и презентации CRM.

Запуск:
    python3 docs/build-docx.py
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOCS_DIR = Path(__file__).resolve().parent

BRAND_PRIMARY = RGBColor(0x0F, 0x3D, 0x91)      # deep blue
BRAND_ACCENT = RGBColor(0xE1, 0x1F, 0x2A)       # red accent
TEXT_MAIN = RGBColor(0x11, 0x18, 0x27)          # near-black
TEXT_MUTED = RGBColor(0x51, 0x5B, 0x6E)         # gray
TABLE_HEADER_BG = "0F3D91"
TABLE_ZEBRA_BG = "F2F5FB"
CODE_BG = "F5F5F7"

FONT_MAIN = "Calibri"
FONT_MONO = "Consolas"


# ---------------------------------------------------------------------------
# Low-level XML helpers
# ---------------------------------------------------------------------------


def _shade(cell_or_paragraph, fill_hex: str) -> None:
    tc_pr = cell_or_paragraph._tc.get_or_add_tcPr() if hasattr(cell_or_paragraph, "_tc") else None
    if tc_pr is None:
        return
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def _paragraph_shade(paragraph, fill_hex: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    p_pr.append(shd)


def _set_cell_borders(cell, color_hex: str = "D0D5DD", size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color_hex)
        tc_borders.append(el)
    tc_pr.append(tc_borders)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


# ---------------------------------------------------------------------------
# Styles
# ---------------------------------------------------------------------------


def configure_styles(doc: Document) -> None:
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = FONT_MAIN
    normal.font.size = Pt(11)
    normal.font.color.rgb = TEXT_MAIN
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, bold, before, after in [
        ("Heading 1", 22, BRAND_PRIMARY, True, 18, 8),
        ("Heading 2", 16, BRAND_PRIMARY, True, 14, 6),
        ("Heading 3", 13, BRAND_PRIMARY, True, 10, 4),
        ("Heading 4", 12, TEXT_MAIN, True, 8, 3),
    ]:
        style = styles[name]
        style.font.name = FONT_MAIN
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.0)


def add_footer(doc: Document, title: str) -> None:
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{title}   ·   стр. ")
        run.font.size = Pt(9)
        run.font.color.rgb = TEXT_MUTED
        _add_page_number(p)
        for run in p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = TEXT_MUTED


# ---------------------------------------------------------------------------
# Inline markdown parsing (bold, italic, code)
# ---------------------------------------------------------------------------

INLINE_PATTERN = re.compile(
    r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*]+\*|_[^_]+_|`[^`]+`)"
)


def add_inline(paragraph, text: str, *, base_bold: bool = False) -> None:
    """Добавить текст с базовым markdown-инлайн-форматированием."""
    text = text.replace("\\(", "(").replace("\\)", ")")
    parts = INLINE_PATTERN.split(text)
    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = FONT_MAIN
        run.font.size = Pt(11)
        if part.startswith("**") and part.endswith("**"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("__") and part.endswith("__"):
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            run.text = part[1:-1]
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run.text = part[1:-1]
            run.font.name = FONT_MONO
            run.font.color.rgb = BRAND_ACCENT
        else:
            run.text = part
        if base_bold:
            run.bold = True


# ---------------------------------------------------------------------------
# Markdown parser (subset, enough for our docs)
# ---------------------------------------------------------------------------


@dataclass
class Block:
    kind: str  # "h1" | "h2" | "h3" | "h4" | "p" | "ul" | "ol" | "hr" | "table" | "code" | "quote"
    payload: object


def parse_markdown(md: str) -> list[Block]:
    lines = md.replace("\r\n", "\n").split("\n")
    blocks: list[Block] = []
    i = 0
    n = len(lines)

    def flush_paragraph(buf: list[str]) -> None:
        if not buf:
            return
        text = " ".join(part.strip() for part in buf).strip()
        if text:
            blocks.append(Block("p", text))
        buf.clear()

    paragraph_buf: list[str] = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # Fenced code block
        if stripped.startswith("```"):
            flush_paragraph(paragraph_buf)
            code_lines: list[str] = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append(Block("code", "\n".join(code_lines)))
            continue

        # Horizontal rule
        if re.match(r"^\s*(---|\*\*\*|___)\s*$", line):
            flush_paragraph(paragraph_buf)
            blocks.append(Block("hr", None))
            i += 1
            continue

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush_paragraph(paragraph_buf)
            level = len(m.group(1))
            blocks.append(Block(f"h{min(level, 4)}", m.group(2).strip()))
            i += 1
            continue

        # Table (must have separator on next line)
        if "|" in line and i + 1 < n and re.match(r"^\s*\|?\s*:?-{2,}", lines[i + 1]):
            flush_paragraph(paragraph_buf)
            header = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2  # skip separator
            rows: list[list[str]] = []
            while i < n and "|" in lines[i] and lines[i].strip():
                row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row)
                i += 1
            blocks.append(Block("table", (header, rows)))
            continue

        # Blockquote
        if stripped.startswith("> "):
            flush_paragraph(paragraph_buf)
            quote_lines: list[str] = []
            while i < n and lines[i].strip().startswith("> "):
                quote_lines.append(lines[i].strip()[2:])
                i += 1
            blocks.append(Block("quote", " ".join(quote_lines).strip()))
            continue

        # Unordered list
        if re.match(r"^\s*[-*]\s+", line):
            flush_paragraph(paragraph_buf)
            items: list[str] = []
            while i < n and re.match(r"^\s*[-*]\s+", lines[i]):
                item = re.sub(r"^\s*[-*]\s+", "", lines[i]).strip()
                j = i + 1
                while j < n and lines[j].startswith("  ") and not re.match(r"^\s*[-*]\s+", lines[j]) and lines[j].strip():
                    item += " " + lines[j].strip()
                    j += 1
                items.append(item)
                i = j
            blocks.append(Block("ul", items))
            continue

        # Ordered list
        if re.match(r"^\s*\d+\.\s+", line):
            flush_paragraph(paragraph_buf)
            items = []
            while i < n and re.match(r"^\s*\d+\.\s+", lines[i]):
                item = re.sub(r"^\s*\d+\.\s+", "", lines[i]).strip()
                j = i + 1
                while j < n and lines[j].startswith("   ") and not re.match(r"^\s*\d+\.\s+", lines[j]) and lines[j].strip():
                    item += " " + lines[j].strip()
                    j += 1
                items.append(item)
                i = j
            blocks.append(Block("ol", items))
            continue

        # Blank line = paragraph boundary
        if not stripped:
            flush_paragraph(paragraph_buf)
            i += 1
            continue

        paragraph_buf.append(stripped)
        i += 1

    flush_paragraph(paragraph_buf)
    return blocks


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------


def render_heading(doc: Document, level: int, text: str) -> None:
    p = doc.add_heading(level=level)
    add_inline(p, text)
    if level == 1:
        p.paragraph_format.space_before = Pt(24)


def render_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    add_inline(p, text)


def render_list(doc: Document, items: Iterable[str], *, ordered: bool) -> None:
    style = "List Number" if ordered else "List Bullet"
    for item in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(3)
        add_inline(p, item)


def render_code(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.4)
    _paragraph_shade(p, CODE_BG)
    run = p.add_run(code)
    run.font.name = FONT_MONO
    run.font.size = Pt(9.5)
    run.font.color.rgb = TEXT_MAIN


def render_hr(doc: Document) -> None:
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "0F3D91")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)


def render_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.right_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    _paragraph_shade(p, "F2F5FB")
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = TEXT_MUTED


def render_table(doc: Document, header: list[str], rows: list[list[str]]) -> None:
    n_cols = len(header)
    table = doc.add_table(rows=1 + len(rows), cols=n_cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row
    hdr = table.rows[0]
    for idx, text in enumerate(header):
        cell = hdr.cells[idx]
        _shade(cell, TABLE_HEADER_BG)
        _set_cell_borders(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].clear()
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(10.5)
        run.font.name = FONT_MAIN

    # Body rows
    for r_idx, row in enumerate(rows):
        tr = table.rows[r_idx + 1]
        is_zebra = r_idx % 2 == 1
        for c_idx in range(n_cols):
            cell = tr.cells[c_idx]
            text = row[c_idx] if c_idx < len(row) else ""
            _set_cell_borders(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if is_zebra:
                _shade(cell, TABLE_ZEBRA_BG)
            cell.paragraphs[0].clear()
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text)
            for run in p.runs:
                run.font.size = Pt(10)


def render_blocks(doc: Document, blocks: list[Block]) -> None:
    for block in blocks:
        if block.kind in ("h1", "h2", "h3", "h4"):
            render_heading(doc, int(block.kind[1:]), block.payload)  # type: ignore[arg-type]
        elif block.kind == "p":
            render_paragraph(doc, block.payload)  # type: ignore[arg-type]
        elif block.kind == "ul":
            render_list(doc, block.payload, ordered=False)  # type: ignore[arg-type]
        elif block.kind == "ol":
            render_list(doc, block.payload, ordered=True)  # type: ignore[arg-type]
        elif block.kind == "hr":
            render_hr(doc)
        elif block.kind == "code":
            render_code(doc, block.payload)  # type: ignore[arg-type]
        elif block.kind == "table":
            header, rows = block.payload  # type: ignore[misc]
            render_table(doc, header, rows)
        elif block.kind == "quote":
            render_quote(doc, block.payload)  # type: ignore[arg-type]


# ---------------------------------------------------------------------------
# Cover
# ---------------------------------------------------------------------------


def add_cover(
    doc: Document,
    *,
    eyebrow: str,
    title: str,
    subtitle: str,
    meta_lines: list[str],
) -> None:
    for _ in range(2):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(eyebrow.upper())
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = BRAND_ACCENT

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(32)
    run.font.color.rgb = BRAND_PRIMARY

    if subtitle:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(24)
        run = p.add_run(subtitle)
        run.font.size = Pt(15)
        run.font.color.rgb = TEXT_MUTED

    # Divider
    render_hr(doc)

    for line in meta_lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line)
        run.font.size = Pt(10.5)
        run.font.color.rgb = TEXT_MAIN

    doc.add_page_break()


# ---------------------------------------------------------------------------
# Post-processing: split TZ cover from main content
# ---------------------------------------------------------------------------


def strip_leading_metadata(md: str, keep_from_heading: str) -> tuple[str, list[str]]:
    """Разбирает header markdown-файла ТЗ и возвращает (тело, meta-строки)."""
    lines = md.replace("\r\n", "\n").split("\n")
    meta: list[str] = []
    body_start = 0
    for idx, line in enumerate(lines):
        if line.strip().startswith(keep_from_heading):
            body_start = idx + 1
            continue
        if body_start and line.strip() == "---":
            body_start = idx + 1
            break
        if body_start and line.strip():
            meta.append(line.strip())
    body = "\n".join(lines[body_start:]).lstrip()
    return body, meta


# ---------------------------------------------------------------------------
# Builders
# ---------------------------------------------------------------------------


def build_tz_docx(md_path: Path, out_path: Path) -> None:
    md = md_path.read_text(encoding="utf-8")
    body, meta = strip_leading_metadata(md, "# Техническое задание")

    doc = Document()
    configure_styles(doc)
    add_cover(
        doc,
        eyebrow="Техническое задание",
        title="CRM «TiSei / Береке ТехСервис»",
        subtitle="Диспетчеризация, ТО, склад, аналитика, мобильное приложение",
        meta_lines=meta or [
            "Версия: 1.0",
            "Дата: август 2026 г.",
            "Заказчик: ТОО «Береке ТехСервис»",
            "Территория эксплуатации: г. Астана, Республика Казахстан",
        ],
    )
    add_footer(doc, "ТЗ. CRM «TiSei / Береке ТехСервис»")

    blocks = parse_markdown(body)
    render_blocks(doc, blocks)

    doc.save(out_path)


def build_presentation_docx(md_path: Path, out_path: Path) -> None:
    md = md_path.read_text(encoding="utf-8")
    lines = md.replace("\r\n", "\n").split("\n")
    body_start = 0
    for idx, line in enumerate(lines):
        if line.strip().startswith("# CRM"):
            body_start = idx + 1
            break
    body = "\n".join(lines[body_start:]).lstrip()

    doc = Document()
    configure_styles(doc)
    add_cover(
        doc,
        eyebrow="Функционал CRM",
        title="CRM «TiSei / Береке ТехСервис»",
        subtitle="Материалы для презентации и коммерческих встреч",
        meta_lines=[
            "20 слайдов + демо-сценарий + ключевые цифры",
            "Формат: тезисы, готовые к переносу в PPTX/Keynote/Figma",
            "Обновлено: август 2026 г.",
        ],
    )
    add_footer(doc, "CRM «TiSei / Береке ТехСервис». Функционал")

    blocks = parse_markdown(body)
    render_blocks(doc, blocks)

    doc.save(out_path)


# ---------------------------------------------------------------------------
# Entry
# ---------------------------------------------------------------------------


def main() -> None:
    tz_md = DOCS_DIR / "TZ_CRM.md"
    pres_md = DOCS_DIR / "CRM_PRESENTATION.md"

    tz_out = DOCS_DIR / "TZ_CRM.docx"
    pres_out = DOCS_DIR / "CRM_PRESENTATION.docx"

    print(f"Building {tz_out.name} from {tz_md.name}…")
    build_tz_docx(tz_md, tz_out)
    print(f"  ✓ {tz_out} ({tz_out.stat().st_size / 1024:.1f} KB)")

    print(f"Building {pres_out.name} from {pres_md.name}…")
    build_presentation_docx(pres_md, pres_out)
    print(f"  ✓ {pres_out} ({pres_out.stat().st_size / 1024:.1f} KB)")


if __name__ == "__main__":
    main()
